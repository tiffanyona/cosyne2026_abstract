"""
tools/extract_pptx_layout.py
-----------------------------
Extract panel positions and sizes from a PowerPoint draft figure.

Workflow
--------
1. In PowerPoint, open the figure slide and rename each shape to its panel key
   (e.g. "fig2_A", "fig2_B") using Home → Select → Selection Pane.
2. Run:
       python tools/extract_pptx_layout.py figures/fig2/fig2.pptx --figure fig2
   or for a specific slide (0-indexed):
       python tools/extract_pptx_layout.py figures/fig2/fig2.pptx --figure fig2 --slide 1
3. The script writes two files alongside the .pptx:
       figures/fig2/layout.csv   — exact mm positions and sizes
       figures/fig2/layout.tex   — LaTeX \\newlength definitions for each panel
4. Recompile the .tex document — no manual fraction arithmetic required.

Output
------
- figures/{figure}/layout.csv  written automatically
- figures/{figure}/layout.tex  written automatically
- Python dict block printed to stdout (for reference)
- Markdown table printed to stdout (for reference)

All values in mm, x/y from the upper-left corner of the bounding box that
encloses all matching shapes (so the first panel starts at x=0, y=0).
"""

import argparse
import csv
import os
import sys

EMU_PER_MM = 36000

_ROMAN = {"1": "I", "2": "II", "3": "III", "4": "IV", "5": "V",
          "6": "VI", "7": "VII", "8": "VIII", "9": "IX", "10": "X"}


def _roman(fig_prefix):
    """'fig2' → 'II'"""
    num = fig_prefix.lstrip("fig").lstrip("0") or "0"
    return _ROMAN.get(num, num.upper())


def _length_name(panel_name, suffix):
    """'fig2_A', 'W' → '\\FigIIAW'"""
    parts = panel_name.split("_")
    roman = _roman(parts[0])
    letter = parts[1].upper()
    return f"\\Fig{roman}{letter}{suffix}"


def extract(pptx_path, figure_prefix, slide_index):
    try:
        from pptx import Presentation
    except ImportError:
        sys.exit("python-pptx is not installed. Run: pip install python-pptx")

    prs = Presentation(pptx_path)
    if slide_index >= len(prs.slides):
        sys.exit(f"Slide {slide_index} not found (presentation has {len(prs.slides)} slide(s)).")
    slide = prs.slides[slide_index]

    def collect(shapes, out):
        for shape in shapes:
            name = shape.name.strip()
            if name.startswith(figure_prefix + "_") or name == figure_prefix:
                out[name] = {
                    "left":   shape.left,
                    "top":    shape.top,
                    "width":  shape.width,
                    "height": shape.height,
                }
            if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
                collect(shape.shapes, out)

    panels = {}
    collect(slide.shapes, panels)

    if not panels:
        sys.exit(
            f'No shapes found with prefix "{figure_prefix}_" on slide {slide_index}.\n'
            f'Available shape names: {[s.name for s in slide.shapes]}'
        )

    # Normalize so the top-left corner of the bounding box = (0, 0)
    min_left = min(v["left"] for v in panels.values())
    min_top  = min(v["top"]  for v in panels.values())

    result = {}
    for name, emu in panels.items():
        result[name] = {
            "x": round((emu["left"]  - min_left) / EMU_PER_MM, 1),
            "y": round((emu["top"]   - min_top)  / EMU_PER_MM, 1),
            "w": round( emu["width"]             / EMU_PER_MM, 1),
            "h": round( emu["height"]            / EMU_PER_MM, 1),
        }

    # Sort by (y, x) — reading order top-to-bottom, left-to-right
    return dict(sorted(result.items(), key=lambda kv: (kv[1]["y"], kv[1]["x"])))


def print_layout_dict(panels, figure_prefix):
    print(f"\n# ---- {figure_prefix.upper()} layout (for reference) ----")
    for name, v in panels.items():
        print(f'    "{name}": {{"x": {v["x"]:6.1f}, "y": {v["y"]:6.1f}, "w": {v["w"]:6.1f}, "h": {v["h"]:5.1f}}},')


def print_table(panels, figure_prefix):
    total_w = max(v["x"] + v["w"] for v in panels.values())

    print(f"\n## {figure_prefix.upper()} panel layout\n")
    print(f"| Panel | x (mm) | y (mm) | w (mm) | h (mm) | w / total |")
    print(f"|-------|--------|--------|--------|--------|-----------|")
    for name, v in panels.items():
        letter = name.split("_")[-1].upper()
        frac = v["w"] / total_w
        print(f"| {letter:<5} | {v['x']:6.1f} | {v['y']:6.1f} | {v['w']:6.1f} | {v['h']:6.1f} | {frac:.3f}     |")

    print(f"\nTotal figure width: {total_w:.1f} mm")


def write_csv(panels, out_path):
    """Write exact panel positions and sizes to a CSV file."""
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    with open(out_path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["panel", "x_mm", "y_mm", "w_mm", "h_mm"])
        writer.writeheader()
        for name, v in panels.items():
            writer.writerow({
                "panel": name,
                "x_mm":  v["x"],
                "y_mm":  v["y"],
                "w_mm":  v["w"],
                "h_mm":  v["h"],
            })
    print(f"wrote {out_path}")


def write_layout_tex(panels, figure_prefix, out_path):
    """Write LaTeX \\newlength definitions for every panel plus derived group widths."""
    roman = _roman(figure_prefix)
    pptx_name = os.path.basename(out_path).replace("layout.tex", f"{figure_prefix}.pptx")

    lines = [
        f"% figures/{figure_prefix}/layout.tex — auto-generated, do not edit by hand",
        f"% Regenerate: python tools/extract_pptx_layout.py figures/{figure_prefix}/{pptx_name} --figure {figure_prefix}",
        "",
        "% Per-panel positions and sizes (exact PowerPoint measurements in mm)",
    ]

    for name, v in sorted(panels.items()):
        letter = name.split("_")[-1].upper()
        x_cmd = f"\\Fig{roman}{letter}X"
        y_cmd = f"\\Fig{roman}{letter}Y"
        w_cmd = f"\\Fig{roman}{letter}W"
        h_cmd = f"\\Fig{roman}{letter}H"
        lines.append(f"\\newlength{{{x_cmd}}}\\setlength{{{x_cmd}}}{{{v['x']}mm}}")
        lines.append(f"\\newlength{{{y_cmd}}}\\setlength{{{y_cmd}}}{{{v['y']}mm}}")
        lines.append(f"\\newlength{{{w_cmd}}}\\setlength{{{w_cmd}}}{{{v['w']}mm}}")
        lines.append(f"\\newlength{{{h_cmd}}}\\setlength{{{h_cmd}}}{{{v['h']}mm}}")

    # Derived group widths computed from panel coordinates
    lines += ["", "% Derived group widths (bounding boxes of panel clusters)"]

    # Right section: width of the widest full-span panel (L or M)
    right_w_panels = [f"{figure_prefix}_L", f"{figure_prefix}_M"]
    right_w_vals = [panels[k]["w"] for k in right_w_panels if k in panels]
    if right_w_vals:
        right_w = max(right_w_vals)
        lines.append(f"\\newlength{{\\Fig{roman}RightW}}\\setlength{{\\Fig{roman}RightW}}{{{right_w}mm}}")

    # Left sub-block: bounding box of B+C columns
    b, c = f"{figure_prefix}_B", f"{figure_prefix}_C"
    if b in panels and c in panels:
        left_sub_w = round(panels[c]["x"] + panels[c]["w"] - panels[b]["x"], 1)
        lines.append(f"\\newlength{{\\Fig{roman}LeftSubW}}\\setlength{{\\Fig{roman}LeftSubW}}{{{left_sub_w}mm}}")

    # Right sub-block: bounding box of F+G columns
    f_key, g = f"{figure_prefix}_F", f"{figure_prefix}_G"
    if f_key in panels and g in panels:
        right_sub_w = round(panels[g]["x"] + panels[g]["w"] - panels[f_key]["x"], 1)
        lines.append(f"\\newlength{{\\Fig{roman}RightSubW}}\\setlength{{\\Fig{roman}RightSubW}}{{{right_sub_w}mm}}")

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    with open(out_path, "w") as f:
        f.write("\n".join(lines) + "\n")
    print(f"wrote {out_path}")


def main():
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("pptx", help="Path to the PowerPoint file")
    parser.add_argument("--figure", required=True, help='Figure prefix, e.g. "fig2"')
    parser.add_argument("--slide", type=int, default=0, help="Slide index (0-based, default 0)")
    parser.add_argument(
        "--out-dir",
        default=None,
        help="Directory for layout.csv and layout.tex (default: same directory as the .pptx)",
    )
    args = parser.parse_args()

    panels = extract(args.pptx, args.figure, args.slide)
    print_layout_dict(panels, args.figure)
    print_table(panels, args.figure)

    out_dir = args.out_dir or os.path.dirname(os.path.abspath(args.pptx))
    write_csv(panels,        os.path.join(out_dir, "layout.csv"))
    write_layout_tex(panels, args.figure, os.path.join(out_dir, "layout.tex"))


if __name__ == "__main__":
    main()
